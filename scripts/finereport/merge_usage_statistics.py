import argparse
import csv
import json
import re
import tempfile
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


FIELD_ALIASES = {
    "report_path": [
        "报表路径",
        "报表模板路径",
        "模板路径",
        "资源路径",
        "viewlet",
        "viewlet路径",
        "cpt路径",
        "cpt",
        "report_path",
        "template_path",
    ],
    "report_name": ["报表名称", "模板名称", "report_name", "name"],
    "user_name": ["使用人", "使用人员", "访问人", "执行人", "用户名", "用户姓名", "用户", "姓名", "user_name", "username"],
    "position": ["岗位", "职位", "职务", "position", "job_title"],
    "department": ["部门", "组织", "所属部门", "department", "dept"],
    "usage_count": ["使用次数", "访问次数", "执行次数", "查询次数", "次数", "usage_count", "count"],
}

EXCLUDED_USAGE_USER_KEYWORDS = ("超管", "admin")


def normalize_header(value):
    return re.sub(r"\s+", "", str(value or "")).lower()


def pick_column(headers, logical_name):
    normalized = {normalize_header(header): index for index, header in enumerate(headers)}
    for alias in FIELD_ALIASES[logical_name]:
        key = normalize_header(alias)
        if key in normalized:
            return normalized[key]
    return None


def header_score(headers):
    columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
    score = sum(1 for value in columns.values() if value is not None)
    if columns["user_name"] is not None:
        score += 3
    if columns["report_path"] is not None or columns["report_name"] is not None:
        score += 3
    if columns["usage_count"] is not None:
        score += 2
    return score


def split_merged_header(value):
    text = cell_text(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[\r\n\t]+", text) if part.strip()]


def detect_table(rows, scan_limit=30):
    candidates = []
    for row_index, row in enumerate(rows[:scan_limit]):
        headers = [cell_text(value) for value in row]
        score = header_score(headers)
        if len(headers) == 1:
            split_headers = split_merged_header(headers[0])
            if len(split_headers) > 1:
                split_score = header_score(split_headers)
                if split_score > score:
                    headers = split_headers
                    score = split_score
        candidates.append((score, row_index, headers))
    score, row_index, headers = max(candidates, key=lambda item: (item[0], -item[1]))
    if score == 0:
        preview = [{"row": index + 1, "values": [cell_text(value) for value in row]} for index, row in enumerate(rows[:10])]
        raise ValueError("could not detect usage header row: " + json.dumps(preview, ensure_ascii=False))
    return row_index, headers


def cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def read_xlsx(path):
    return workbook_rows_to_records(read_xlsx_rows(path))


def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    if not rows:
        return []
    header_row_index, headers = detect_table(rows)
    return rows_to_records(headers, rows[header_row_index + 1 :])


def inspect_usage_source(source, save_download=None):
    if re.match(r"^https?://", source, re.I):
        source_path = download_usage_file(source, save_download)
    else:
        source_path = Path(source)
    suffix = source_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return inspect_workbook_rows(source_path, read_xlsx_rows(source_path))
    elif suffix == ".csv":
        with open(source_path, "r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    else:
        raise ValueError(f"unsupported usage file type: {source_path.suffix}; use .xlsx or .csv")
    header_row_index, headers = detect_table(rows)
    columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
    return {
        "source": str(source_path),
        "detected_header_row": header_row_index + 1,
        "headers": headers,
        "recognized_columns": {
            name: headers[index] if index is not None and index < len(headers) else None
            for name, index in columns.items()
        },
        "preview_rows": [
            [cell_text(value) for value in row]
            for row in rows[header_row_index + 1 : header_row_index + 6]
        ],
    }


def inspect_workbook_rows(source_path, workbook_rows):
    sheets = []
    for sheet_name, rows in workbook_rows:
        if not rows:
            sheets.append({"sheet": sheet_name, "rows": 0, "columns": 0, "detected_header_row": None})
            continue
        try:
            header_row_index, headers = detect_table(rows)
            columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
            recognized = {
                name: headers[index] if index is not None and index < len(headers) else None
                for name, index in columns.items()
            }
        except ValueError:
            header_row_index, headers, recognized = None, [], {}
        sheets.append({
            "sheet": sheet_name,
            "rows": len(rows),
            "columns": max((len(row) for row in rows), default=0),
            "detected_header_row": header_row_index + 1 if header_row_index is not None else None,
            "headers": headers,
            "recognized_columns": recognized,
            "preview_rows": [
                [cell_text(value) for value in row]
                for row in rows[(header_row_index + 1 if header_row_index is not None else 0) : (header_row_index + 6 if header_row_index is not None else 5)]
            ],
        })
    return {"source": str(source_path), "sheets": sheets}


def workbook_rows_to_records(workbook_rows):
    usage_records = []
    people = {}
    diagnostics = []
    for sheet_name, rows in workbook_rows:
        if not rows:
            continue
        try:
            header_row_index, headers = detect_table(rows)
        except ValueError as error:
            diagnostics.append({"sheet": sheet_name, "error": str(error)})
            continue
        columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
        data_rows = rows[header_row_index + 1 :]
        if columns["user_name"] is not None and (columns["report_path"] is not None or columns["report_name"] is not None):
            usage_records.extend(rows_to_records(headers, data_rows))
        elif columns["user_name"] is not None and (columns["position"] is not None or columns["department"] is not None):
            people.update(rows_to_people(headers, data_rows))
        else:
            diagnostics.append({"sheet": sheet_name, "headers": headers, "error": "sheet is neither usage nor people metadata"})

    if people:
        for record in usage_records:
            person = people.get(record.get("user_name", ""))
            if person:
                record["position"] = record.get("position") or person.get("position", "")
                record["department"] = record.get("department") or person.get("department", "")
    if usage_records:
        return usage_records
    raise ValueError("no usage detail rows found in workbook: " + json.dumps(diagnostics, ensure_ascii=False))


def read_xlsx_rows(path):
    path = Path(path)
    if zipfile.is_zipfile(path):
        rows = read_xlsx_rows_from_xml(path)
        if rows:
            return rows
    workbook = load_workbook(path, read_only=True, data_only=True)
    return [(sheet.title, [[cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]


def read_xlsx_rows_from_xml(path):
    namespaces = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
        "office_rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.attrib["Id"]: rel.attrib["Target"]
            for rel in workbook_rels.findall("rel:Relationship", namespaces)
        }
        result = []
        for sheet in workbook.findall("main:sheets/main:sheet", namespaces):
            sheet_name = sheet.attrib.get("name", "")
            rel_id = sheet.attrib.get(f"{{{namespaces['office_rel']}}}id")
            target = rel_targets.get(rel_id, "")
            sheet_path = "xl/" + target.lstrip("/")
            if sheet_path not in archive.namelist():
                continue
            result.append((sheet_name, worksheet_xml_to_rows(archive.read(sheet_path))))
        return result


def worksheet_xml_to_rows(content):
    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(content)
    rows = defaultdict(dict)
    for cell in root.findall(".//main:c", namespace):
        reference = cell.attrib.get("r", "")
        if not reference:
            continue
        rows[row_number(reference)][column_number(reference)] = cell_value(cell, namespace)
    if not rows:
        return []
    width = max((max(columns.keys()) for columns in rows.values() if columns), default=-1) + 1
    return [[rows[row_index].get(column_index, "") for column_index in range(width)] for row_index in range(max(rows.keys()) + 1)]


def column_number(reference):
    letters = "".join(character for character in reference if character.isalpha())
    number = 0
    for character in letters:
        number = number * 26 + ord(character.upper()) - 64
    return number - 1


def row_number(reference):
    return int("".join(character for character in reference if character.isdigit())) - 1


def cell_value(cell, namespace):
    value = cell.find("main:v", namespace)
    if value is not None:
        return cell_text(value.text)
    return cell_text("".join(text.text or "" for text in cell.findall(".//main:t", namespace)))


def rows_to_people(headers, rows):
    columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
    people = {}
    for row in rows:
        if not any(cell_text(value) for value in row):
            continue
        user_name = normalize_user_name(row[columns["user_name"]]) if columns["user_name"] is not None and columns["user_name"] < len(row) else ""
        if not user_name:
            continue
        people[user_name] = {
            "position": cell_text(row[columns["position"]]) if columns["position"] is not None and columns["position"] < len(row) else "",
            "department": cell_text(row[columns["department"]]) if columns["department"] is not None and columns["department"] < len(row) else "",
        }
    return people


def rows_to_records(headers, rows):
    columns = {name: pick_column(headers, name) for name in FIELD_ALIASES}
    if columns["user_name"] is None:
        raise ValueError("usage file must include a user column, such as 使用人/用户名/用户")
    if columns["report_path"] is None and columns["report_name"] is None:
        raise ValueError("usage file must include 报表路径/viewlet or 报表名称")

    records = []
    for row in rows:
        if not any(cell_text(value) for value in row):
            continue
        record = {}
        for name, index in columns.items():
            record[name] = cell_text(row[index]) if index is not None and index < len(row) else ""
        raw_user_name = record.get("user_name", "")
        normalized_user_name = normalize_user_name(raw_user_name)
        if is_excluded_usage_user(raw_user_name) or is_excluded_usage_user(normalized_user_name):
            continue
        record["user_name"] = normalized_user_name
        record["usage_count"] = parse_count(record.get("usage_count")) or 1
        records.append(record)
    return records


def is_excluded_usage_user(value):
    text = cell_text(value).lower()
    return any(keyword in text for keyword in EXCLUDED_USAGE_USER_KEYWORDS)


def normalize_user_name(value):
    text = cell_text(value)
    match = re.match(r"^(.*?)\((.*?)\)$", text)
    if match and match.group(1).strip():
        return match.group(1).strip()
    return text


def parse_count(value):
    text = cell_text(value).replace(",", "")
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def download_usage_file(url, save_path=None):
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        suffix = guess_suffix(url, response.headers.get("Content-Type", ""))
        if save_path:
            output = Path(save_path)
            if not output.suffix:
                output = output.with_suffix(suffix)
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(response.read())
            return output
        temporary = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        with temporary:
            temporary.write(response.read())
        return Path(temporary.name)


def guess_suffix(url, content_type):
    lower = (content_type or "").lower()
    if "csv" in lower:
        return ".csv"
    if "spreadsheetml" in lower or "xlsx" in lower:
        return ".xlsx"
    path = urllib.parse.urlparse(url).path.lower()
    if path.endswith(".csv"):
        return ".csv"
    if path.endswith(".xlsx"):
        return ".xlsx"
    return ".xlsx"


def read_usage(source, save_download=None):
    if re.match(r"^https?://", source, re.I):
        source_path = download_usage_file(source, save_download)
    else:
        source_path = Path(source)
    suffix = source_path.suffix.lower()
    if suffix == ".csv":
        return read_csv(source_path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(source_path)
    raise ValueError(f"unsupported usage file type: {source_path.suffix}; use .xlsx or .csv")


def normalize_report_key(value):
    text = urllib.parse.unquote(urllib.parse.unquote(cell_text(value))).replace("\\", "/")
    text = re.sub(r"/+", "/", text).strip().lower()
    return text


def report_keys_from_json(path, payload):
    report = payload.get("report", {})
    keys = set()
    for value in [report.get("source_cpt_path"), report.get("name"), path.stem]:
        normalized = normalize_report_key(value)
        if normalized:
            keys.add(normalized)
            keys.add(Path(normalized).name)
    source_path = normalize_report_key(report.get("source_cpt_path"))
    marker = "/lr/"
    if marker in source_path:
        keys.add(source_path.split(marker, 1)[1])
    return keys


def record_keys(record):
    keys = set()
    for value in [record.get("report_path"), record.get("report_name")]:
        normalized = normalize_report_key(value)
        if normalized:
            keys.add(normalized)
            keys.add(Path(normalized).name)
    return keys


def build_usage_index(records):
    usage = defaultdict(lambda: defaultdict(lambda: {"usage_count": 0, "position": "", "department": ""}))
    totals = defaultdict(int)
    for record in records:
        keys = record_keys(record)
        user_name = record.get("user_name", "")
        if not keys or not user_name:
            continue
        count = record["usage_count"]
        for key in keys:
            totals[key] += count
            user = usage[key][user_name]
            user["usage_count"] += count
            user["position"] = user["position"] or record.get("position", "")
            user["department"] = user["department"] or record.get("department", "")
    return usage, totals


def usage_for_report(json_path, payload, usage_index, total_index):
    merged_users = defaultdict(lambda: {"usage_count": 0, "position": "", "department": ""})
    total = 0
    for key in report_keys_from_json(json_path, payload):
        total += total_index.get(key, 0)
        for user_name, user in usage_index.get(key, {}).items():
            merged = merged_users[user_name]
            merged["usage_count"] += user["usage_count"]
            merged["position"] = merged["position"] or user["position"]
            merged["department"] = merged["department"] or user["department"]
    users = [
        {
            "user_name": user_name,
            "position": data["position"],
            "department": data["department"],
            "usage_count": data["usage_count"],
        }
        for user_name, data in merged_users.items()
    ]
    users.sort(key=lambda item: (-item["usage_count"], item["user_name"]))
    return {"total_usage_count": total, "users": users}


def write_json(path, payload):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def merge_usage(usage_source, parsed_dir, dry_run=False, save_download=None):
    records = read_usage(usage_source, save_download)
    usage_index, total_index = build_usage_index(records)
    updated = 0
    matched = 0
    for json_path in sorted(Path(parsed_dir).rglob("*.json")):
        payload = json.loads(json_path.read_text(encoding="utf-8"))
        report_usage = usage_for_report(json_path, payload, usage_index, total_index)
        if report_usage["total_usage_count"]:
            matched += 1
        payload["report_usage"] = report_usage
        updated += 1
        if not dry_run:
            write_json(json_path, payload)
    return {"usage_rows": len(records), "json_files": updated, "matched_reports": matched}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("usage_source", help="usage Excel/CSV path or FineReport excel URL")
    parser.add_argument("parsed_json_dir", nargs="?", help="directory containing parsed report JSON files")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--inspect", action="store_true", help="print detected headers and sample rows")
    parser.add_argument("--save-download", help="save downloaded URL content to this Excel/CSV path")
    args = parser.parse_args()
    if args.inspect:
        print(json.dumps(inspect_usage_source(args.usage_source, args.save_download), ensure_ascii=False, indent=2))
        return
    if not args.parsed_json_dir:
        raise SystemExit("parsed_json_dir is required unless --inspect is used")
    summary = merge_usage(args.usage_source, args.parsed_json_dir, args.dry_run, args.save_download)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
